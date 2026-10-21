import hashlib
import io
import json
import math
import os
import re
import uuid
import warnings
from datetime import datetime, timezone
from contextlib import contextmanager
from pathlib import Path
import pytest
from docx import Document
from sqlalchemy import create_engine, text
from sqlalchemy.engine import make_url
from werkzeug.datastructures import FileStorage
from app import create_app
from app.extensions import db
from app.services.chunking import assign_ids


PHASE35_TEST_PORT = 55439
PHASE35_TEST_ROLE = "rag_phase35_test"


def verify_phase35_test_target(admin_url, engine_factory=create_engine):
    """Fail closed unless PostgreSQL identifies as this repo's private test cluster."""
    parsed = make_url(admin_url)
    expected_data_dir = (Path(__file__).resolve().parents[1] / "instance" / "pg-phase35").resolve()
    if (parsed.get_backend_name() != "postgresql" or parsed.host not in ("127.0.0.1", "localhost", "::1")
            or parsed.port != PHASE35_TEST_PORT or parsed.database != "postgres"
            or parsed.username != PHASE35_TEST_ROLE):
        raise RuntimeError("PostgreSQL target is not the designated localhost Phase 3.5 test admin role/port.")
    identity_file = expected_data_dir / "phase35-test-identity.txt"
    if not identity_file.is_file():
        raise RuntimeError("The isolated PostgreSQL cluster has no private identity marker.")
    cluster_identity = identity_file.read_text(encoding="utf8").strip()
    expected_identity = "phase35-test:" + hashlib.sha256(
        str(expected_data_dir).casefold().encode("utf8")).hexdigest()
    if cluster_identity != expected_identity:
        raise RuntimeError("The private PostgreSQL cluster identity marker does not match the expected data directory.")
    engine = engine_factory(parsed)
    try:
        with engine.connect() as connection:
            port = int(connection.exec_driver_sql("SHOW port").scalar_one())
            version = connection.exec_driver_sql("SHOW server_version").scalar_one()
            user, database = connection.exec_driver_sql("SELECT current_user, current_database()").one()
            observed_identity = connection.exec_driver_sql(
                "SELECT shobj_description(oid, 'pg_database') FROM pg_database WHERE datname = current_database()"
            ).scalar_one_or_none()
            createdb, superuser, createrole = connection.exec_driver_sql(
                "SELECT rolcreatedb, rolsuper, rolcreaterole FROM pg_roles WHERE rolname = current_user"
            ).one()
            if (observed_identity != cluster_identity or port != PHASE35_TEST_PORT or database != "postgres"
                    or user != PHASE35_TEST_ROLE or not createdb or superuser or createrole):
                raise RuntimeError("Connected PostgreSQL cluster identity or test-role privileges do not match the isolated target.")
    finally:
        engine.dispose()
    artifact = {
        "purpose": "private Phase 3.5 disposable integration-test environment",
        "verified_at_utc": datetime.now(timezone.utc).isoformat(),
        "host": parsed.host,
        "port": port,
        "control_database": database,
        "admin_role": user,
        "role_privileges": {"createdb": bool(createdb), "superuser": bool(superuser), "createrole": bool(createrole)},
        "data_directory": str(expected_data_dir),
        "database_identity_marker_verified": True,
        "server_version": version,
        "test_databases": "rag_test_<random-uuid>; created and dropped by ownership-marked fixture",
        "chroma_isolation": "per-test pytest tmp_path/DATA_DIR/chroma",
        "credentials_recorded": False,
    }
    artifact_path = expected_data_dir.parent / "test-results" / "phase35" / "environment.json"
    artifact_path.parent.mkdir(parents=True, exist_ok=True)
    artifact_path.write_text(json.dumps(artifact, ensure_ascii=False, indent=2) + "\n", encoding="utf8")
    return artifact


class FakeOpenAI:
    """Deterministic OpenAI test double; not evidence of retrieval quality."""
    def __init__(self, deadline=None):
        self.attempts, self.usage = 0, {}
    def count(self, text, model=None):
        return max(1, math.ceil(len(text.encode()) / 3))
    def fit_chunks(self, chunks, model, version_id, config):
        return assign_ids(chunks, version_id, config)
    def embed(self, texts, task, model=None, dimensions=None):
        from flask import current_app
        from app.services.vectors import normalize_vector
        dims = dimensions or current_app.config["OPENAI_EMBEDDING_DIMENSIONS"]
        vectors = []
        for item in texts:
            vector = [0.] * dims
            for word in re.findall(r"\w+", item.casefold()):
                vector[int(hashlib.sha256(word.encode()).hexdigest(), 16) % dims] += 1
            if not any(vector):
                vector[0] = 1
            vectors.append(normalize_vector(vector, dims))
        return vectors
    def generate(self, prompt, system):
        return json.dumps(dict(status="answered", answer="Синтетик журмын дагуу таван ажлын өдөр байна. [S1]", sources=["S1"]), ensure_ascii=False)


@pytest.fixture(autouse=True)
def isolate_secrets_and_dotenv(monkeypatch):
    """No developer credentials or .env values enter any test app."""
    import app as application
    monkeypatch.setattr(application, "load_dotenv", lambda *args, **kwargs: False)
    for name in ("OPENAI_API_KEY",):
        monkeypatch.delenv(name, raising=False)


@contextmanager
def disposable_database(admin_url, engine_factory=create_engine):
    """Create a fresh UUID database; ownership, not a name prefix, permits cleanup."""
    parsed = make_url(admin_url)
    if parsed.get_backend_name() != "postgresql":
        raise ValueError("TEST_DATABASE_ADMIN_URL must use PostgreSQL")
    token = uuid.uuid4().hex
    name, marker = "rag_test_" + token, "pytest-owned:" + token
    admin = engine_factory(parsed, isolation_level="AUTOCOMMIT")
    owned = False
    try:
        with admin.connect() as connection:
            # Both SQL literals below contain only a locally generated UUID.
            connection.execute(text(f'CREATE DATABASE "{name}"'))
            connection.execute(text(f"COMMENT ON DATABASE \"{name}\" IS '{marker}'"))
            owned = True
        yield parsed.set(database=name).render_as_string(hide_password=False)
    finally:
        if owned:
            try:
                with admin.connect() as connection:
                    actual = connection.scalar(text("SELECT shobj_description(oid, 'pg_database') "
                        "FROM pg_database WHERE datname = :name"), {"name": name})
                    if actual != marker:
                        raise RuntimeError("Database ownership changed; refusing cleanup")
                    connection.execute(text(f'DROP DATABASE "{name}"'))
            except Exception as exc:
                warnings.warn(f"Disposable database cleanup deferred ({type(exc).__name__}); no forced cleanup performed")
        admin.dispose()


@pytest.fixture
def app(tmp_path):
    admin_url = os.getenv("TEST_DATABASE_ADMIN_URL")
    if not admin_url:
        pytest.skip("Integration isolation unavailable: set TEST_DATABASE_ADMIN_URL to a PostgreSQL admin connection with CREATEDB")
    verify_phase35_test_target(admin_url)
    # Entering the context must succeed before any app/schema is initialized.
    database = disposable_database(admin_url)
    try:
        url = database.__enter__()
    except Exception as exc:
        raise RuntimeError(f"Cannot provision an owned disposable PostgreSQL database ({type(exc).__name__}).") from exc
    try:
        application = create_app(dict(TESTING=True, DATA_DIR=tmp_path, SECRET_KEY="test-secret", SQLALCHEMY_DATABASE_URI=url,
            EMBEDDING_PROVIDER="openai", EMBEDDING_MODEL="text-embedding-3-small", GENERATION_MODEL="gpt-test",
            OPENAI_EMBEDDING_MODEL="text-embedding-3-small", OPENAI_GENERATION_MODEL="gpt-test",
            OPENAI_EMBEDDING_DIMENSIONS=128,
            OPENAI_API_KEY="", OPENAI_FACTORY=FakeOpenAI,
            EMBEDDING_DIMENSIONS=128, TRUSTED_HOSTS=["localhost", "127.0.0.1"]))
        with application.app_context():
            try:
                from flask_migrate import upgrade
                upgrade(directory=str(Path(__file__).resolve().parents[1] / "migrations"))
                yield application
            finally:
                db.session.remove()
                db.engine.dispose()
    finally:
        database.__exit__(None, None, None)


@pytest.fixture
def client(app):
    client = app.test_client()
    client.get("/chat")
    return client


def csrf(client):
    with client.session_transaction() as session:
        return session["csrf"]


def docx_file(text="Request leave five working days in advance.", table=False):
    buffer = io.BytesIO()
    document = Document()
    document.add_heading("Leave", 1)
    document.add_paragraph(text)
    if table:
        row = document.add_table(rows=1, cols=2).rows[0]
        row.cells[0].text, row.cells[1].text = "Days", "5"
    document.save(buffer)
    buffer.seek(0)
    return FileStorage(buffer, filename="policy.docx")


@pytest.fixture
def ready(app):
    from app.services.ingestion import upload, process, approve
    version, duplicate = upload(docx_file(), "Leave policy", "en")
    process(version.id)
    approve(version.id)
    return process(version.id)
