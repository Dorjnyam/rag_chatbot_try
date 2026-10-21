import io
import json
import logging
from contextlib import nullcontext
from pathlib import Path
from types import SimpleNamespace

from docx import Document
from PIL import Image, ImageDraw
from pypdf import PdfWriter
from werkzeug.datastructures import FileStorage


def _parser_config():
    return {
        "PDF_MAX_PAGES": 100, "DOCX_MAX_EXPANDED": 100 * 1024 * 1024,
        "OCR_MIN_CHARACTERS": 40, "OCR_TIMEOUT": 5, "CONVERT_TIMEOUT": 5,
        "PROCESS_TIMEOUT": 30, "TESSERACT_CMD": "tesseract", "TESSDATA_PREFIX": "",
        "OCR_LANGUAGES": "eng", "LIBREOFFICE_CMD": "soffice",
    }


def _encrypted_pdf(path):
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    writer.encrypt("secret")
    with path.open("wb") as handle:
        writer.write(handle)


def _partial_docx(path, image_path):
    doc = Document()
    doc.add_heading("Журам", level=1)
    doc.add_paragraph("Ажилтан хүсэлтээ урьдчилан өгнө.")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text, table.cell(0, 1).text = "Төрөл", "Хугацаа"
    table.cell(1, 0).text, table.cell(1, 1).text = "Чөлөө", "5 өдөр"
    doc.add_picture(str(image_path))
    doc.save(path)


def test_load_benchmark_records_failures_quality_and_partial_extraction(tmp_path):
    from app.load_benchmark import run_load_benchmark

    malformed_docx = tmp_path / "malformed.docx"
    malformed_docx.write_bytes(b"PK this is not a zip")
    unsupported = tmp_path / "unsupported.txt"
    unsupported.write_text("not a supported format", encoding="utf8")
    empty_docx = tmp_path / "empty.docx"
    Document().save(empty_docx)
    encrypted = tmp_path / "encrypted.pdf"
    _encrypted_pdf(encrypted)
    malformed_pdf = tmp_path / "malformed.pdf"
    malformed_pdf.write_bytes(b"%PDF-1.4\nnot a valid PDF")

    image_path = tmp_path / "partial.png"
    image = Image.new("RGB", (320, 60), "white")
    ImageDraw.Draw(image).text((4, 15), "IMAGE_ONLY_CLAUSE", fill="black")
    image.save(image_path)
    partial = tmp_path / "partial.docx"
    _partial_docx(partial, image_path)
    reference = tmp_path / "partial-reference.txt"
    reference.write_text("Журам\nАжилтан хүсэлтээ урьдчилан өгнө.\nТөрөл | Хугацаа\nЧөлөө | 5 өдөр\nIMAGE_ONLY_CLAUSE", encoding="utf8")

    manifest = tmp_path / "manifest.json"
    manifest.write_text(json.dumps({"schema_version": 1, "name": "edge-cases", "cases": [
        {"id": "bad-docx", "file": "malformed.docx"},
        {"id": "unsupported", "file": "unsupported.txt"},
        {"id": "empty", "file": "empty.docx"},
        {"id": "encrypted", "file": "encrypted.pdf"},
        {"id": "bad-pdf", "file": "malformed.pdf"},
        {"id": "partial", "file": "partial.docx", "ground_truth": {
            "text_file": "partial-reference.txt", "headings": ["Журам"],
            "clauses": ["Ажилтан хүсэлтээ урьдчилан өгнө."],
            "table_texts": ["Чөлөө | 5 өдөр"], "table_blocks": 1}},
    ]}, ensure_ascii=False), encoding="utf8")

    report = run_load_benchmark(manifest, tmp_path / "reports", _parser_config())
    rows = {row["case_id"]: row for row in report["records"]}
    assert rows["bad-docx"]["error_category"] == "malformed_docx"
    assert rows["unsupported"]["error_category"] == "unsupported_format"
    assert rows["empty"]["error_category"] == "empty_document"
    assert rows["encrypted"]["error_category"] == "encrypted"
    assert rows["bad-pdf"]["error_category"] == "malformed_pdf"
    assert rows["partial"]["status"] == "success"
    assert rows["partial"]["heading_recall"] == 1.0
    assert rows["partial"]["table_blocks"] == 1
    assert rows["partial"]["table_block_match"] is True
    assert rows["partial"]["cer"] > 0  # The image-only clause is intentionally not OCRed in DOCX.
    assert report["production_mutations"] is False
    assert sorted(path.suffix for path in (tmp_path / "reports").iterdir()) == [".csv", ".json", ".md"]


def test_upload_duplicate_is_explicitly_deduplicated_without_database(tmp_path, monkeypatch):
    from flask import Flask
    from app.services import ingestion
    from app.services.storage import fingerprint

    content = b"%PDF-synthetic duplicate source"
    match = SimpleNamespace(id="version-1", document_id="document-1", config_hash=fingerprint({}))
    session = SimpleNamespace(scalars=lambda *_: SimpleNamespace(all=lambda: [match]))
    monkeypatch.setattr(ingestion, "db", SimpleNamespace(session=session))
    monkeypatch.setattr(ingestion, "index_config", lambda: {})
    monkeypatch.setattr(ingestion, "mutation_lock", lambda: nullcontext())
    app = Flask(__name__)
    app.config.update(DATA_DIR=tmp_path, MAX_CONTENT_LENGTH=1024)
    captured = io.StringIO()
    handler = logging.StreamHandler(captured)
    app.logger.setLevel(logging.INFO)
    app.logger.addHandler(handler)
    upload = FileStorage(io.BytesIO(content), filename="duplicate.pdf")
    with app.app_context():
        version, duplicate = ingestion.upload(upload, "Duplicate", "en")
    app.logger.removeHandler(handler)
    assert version is match
    assert duplicate is True
    event = next(line for line in captured.getvalue().splitlines() if "rag_upload {" in line)
    fields = json.loads(event.split("rag_upload ", 1)[1])
    assert fields["outcome"] == "duplicate"
    assert fields["document_id"] == "document-1" and fields["version_id"] == "version-1"
    assert fields["run_id"] and "synthetic duplicate source" not in event


def test_upload_mime_mismatch_and_size_overflow_are_observable(tmp_path):
    import pytest
    from flask import Flask
    from app.services import ingestion

    app = Flask(__name__)
    app.config.update(DATA_DIR=tmp_path, MAX_CONTENT_LENGTH=8)
    captured = io.StringIO()
    handler = logging.StreamHandler(captured)
    app.logger.setLevel(logging.INFO)
    app.logger.addHandler(handler)
    with app.app_context():
        wrong_mime = FileStorage(io.BytesIO(b"%PDF-1.4"), filename="policy.pdf", content_type="text/plain")
        with pytest.raises(ValueError, match="type does not match"):
            ingestion.upload(wrong_mime, "Policy", "en")
        too_big = FileStorage(io.BytesIO(b"%PDF-1.4\n"), filename="large.pdf")
        with pytest.raises(ValueError, match="at most 8 bytes"):
            ingestion.upload(too_big, "Large", "en")
    app.logger.removeHandler(handler)
    events = [json.loads(line.split("rag_upload ", 1)[1]) for line in captured.getvalue().splitlines()]
    assert [event["error_category"] for event in events] == ["mime_mismatch", "size_limit_exceeded"]
