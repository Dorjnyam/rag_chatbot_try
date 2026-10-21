"""Audit sentinels for known information loss at the current parser/chunker boundary."""
from uuid import uuid4

from docx import Document

from app.services.chunking import chunk
from app.services.parsing import parse_direct


def test_docx_keeps_typed_clause_text_but_drops_automatic_list_and_heading_levels(tmp_path):
    source = tmp_path / "structure.docx"
    document = Document()
    document.add_heading("4. Parent", level=1)
    document.add_heading("4.2 Child", level=2)
    document.add_paragraph("5.1. Typed clause number stays in text.")
    document.add_paragraph("Automatically numbered item", style="List Number")
    document.add_paragraph("Nested item", style="List Number 2")
    document.save(source)

    artifact = parse_direct(source, "docx", {"DOCX_MAX_EXPANDED": 1_000_000})
    blocks = artifact["blocks"]
    auto_list = next(block for block in blocks if "Automatically numbered" in block["text"])

    assert [block["block"] for block in blocks] == list(range(1, len(blocks) + 1))
    assert any(block["text"].startswith("5.1.") for block in blocks)
    assert not any(block["text"].startswith("1.") for block in blocks if "Automatically numbered" in block["text"])
    assert auto_list["heading"] == "4.2 Child"
    assert not {"heading_level", "parent_heading", "list_level", "list_number"}.intersection(auto_list)
    assert all("page" not in block for block in blocks)


def test_docx_table_and_chunk_boundary_are_flat_text_with_block_offsets(tmp_path):
    source = tmp_path / "table.docx"
    document = Document()
    document.add_heading("5. Benefits", level=1)
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1))
    table.cell(0, 0).text = "Merged header"
    table.cell(1, 0).text = "Cell A"
    table.cell(1, 1).text = "Cell B"
    document.save(source)

    artifact = parse_direct(source, "docx", {"DOCX_MAX_EXPANDED": 1_000_000})
    table_block = next(block for block in artifact["blocks"] if block["kind"] == "table")
    assert table_block["text"] == "Merged header | Merged header\nCell A | Cell B"
    assert not {"rows", "cells", "merged_cell_map", "source_span"}.intersection(table_block)

    chunks = chunk(artifact["blocks"], {"chunk_size": 1_000, "chunk_overlap": 0}, str(uuid4()))
    metadata = next(item["metadata"] for item in chunks if "Merged header" in item["text"])
    assert {"block_start", "block_end", "char_start", "char_end", "heading"}.issubset(metadata)
    assert not {"table_rows", "clause_id", "source_span", "page_start"}.intersection(metadata)
