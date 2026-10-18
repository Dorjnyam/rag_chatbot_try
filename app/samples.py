"""Synthetic policies and draft labels. No claim of bilingual human review."""
import hashlib
import json
from pathlib import Path
from docx import Document

TOPICS = [
    ("leave", "Leave", "Чөлөө", "Request leave 5 working days before the start date.", "Чөлөөний хүсэлтийг эхлэх өдрөөс 5 ажлын өдрийн өмнө өгнө.", "The team lead approves leave requests.", "Багийн ахлагч чөлөөний хүсэлтийг батална.", "How far in advance must I request leave?", "Чөлөөний хүсэлтийг хэд хоногийн өмнө өгөх вэ?", "Who approves leave requests?", "Чөлөөний хүсэлтийг хэн батлах вэ?"),
    ("travel", "Travel", "Томилолт", "The daily meal allowance is 45000 MNT.", "Хоолны өдрийн тэтгэмж 45000 төгрөг байна.", "Submit travel receipts within 7 calendar days after return.", "Томилолтоос ирснээс хойш 7 хуанлийн хоногт баримтаа өгнө.", "What is the daily meal allowance?", "Хоолны өдрийн тэтгэмж хэд вэ?", "When must travel receipts be submitted?", "Томилолтын баримтыг хэзээ өгөх вэ?"),
    ("remote", "Remote work", "Зайнаас ажиллах", "Employees may work remotely 2 days per week.", "Ажилтан долоо хоногт 2 өдөр зайнаас ажиллаж болно.", "Remote work does not apply during equipment installation.", "Тоног төхөөрөмж суурилуулах үед зайнаас ажиллах нөхцөл үйлчлэхгүй.", "How many remote days are allowed each week?", "Долоо хоногт хэдэн өдөр зайнаас ажиллаж болох вэ?", "When does the remote work policy not apply?", "Зайнаас ажиллах нөхцөл хэзээ үйлчлэхгүй вэ?"),
    ("equipment", "Equipment", "Тоног төхөөрөмж", "Report a lost laptop within 2 hours.", "Зөөврийн компьютер алга болбол 2 цагийн дотор мэдэгдэнэ.", "The equipment register is managed by the IT coordinator.", "Тоног төхөөрөмжийн бүртгэлийг мэдээллийн технологийн зохицуулагч хөтөлнө.", "How soon must I report a lost laptop?", "Компьютер алга болсныг хэдийд мэдэгдэх вэ?", "Who manages the equipment register?", "Тоног төхөөрөмжийн бүртгэлийг хэн хөтөлдөг вэ?"),
    ("training", "Training", "Сургалт", "The annual training allowance is 300000 MNT.", "Жилийн сургалтын тэтгэмж 300000 төгрөг байна.", "Training expenses require approval before enrollment.", "Сургалтад бүртгүүлэхээс өмнө зардлаа батлуулна.", "What is the annual training allowance?", "Жилийн сургалтын тэтгэмж хэд вэ?", "When must training expenses be approved?", "Сургалтын зардлыг хэзээ батлуулах вэ?"),
    ("visitors", "Visitors", "Зочин", "Visitor access ends at 18:00.", "Зочны нэвтрэх эрх 18:00 цагт дуусна.", "Visitors must remain with their host in the laboratory.", "Зочин лабораторид хүлээн авагчтайгаа хамт байна.", "When does visitor access end?", "Зочны нэвтрэх эрх хэдэн цагт дуусах вэ?", "Can visitors enter the laboratory alone?", "Зочин лабораторид ганцаараа орж болох уу?"),
    ("expenses", "Expenses", "Зардал", "Expense claims require an itemized receipt.", "Зардал нэхэмжлэхэд задаргаатай баримт шаардлагатай.", "Reimbursements are processed every Friday.", "Зардлын нөхөн төлбөрийг баасан гараг бүр боловсруулна.", "What evidence is needed for an expense claim?", "Зардал нэхэмжлэхэд ямар баримт шаардлагатай вэ?", "When are reimbursements processed?", "Нөхөн төлбөрийг хэзээ боловсруулдаг вэ?"),
    ("safety", "Safety", "Аюулгүй байдал", "The assembly point is the north gate.", "Цугларах цэг нь хойд хаалга байна.", "Do not use elevators during a fire alarm.", "Галын дохиоллын үед цахилгаан шат ашиглаж болохгүй.", "Where is the emergency assembly point?", "Онцгой үед хаана цугларах вэ?", "Can elevators be used during a fire alarm?", "Галын дохиоллын үед цахилгаан шат ашиглаж болох уу?"),
]


def richer_case(index, language, evidence_language, sources):
    """Table and multi-evidence cases remain draft labels until human review."""
    def evidence(topic, fact):
        return dict(sources[(topic[0], evidence_language)], unit_id=f"{topic[0]}-{fact}", block=4 + fact,
            excerpt=topic[3 + fact * 2 + (evidence_language == "mn")])
    if index == 12:
        return dict(question="Сургалтын журмын код юу вэ?" if language == "mn" else "What is the training policy code?",
            expected_answer="TRAINING-01", required_claims=["TRAINING-01"],
            evidence=[dict(sources[("training", evidence_language)], unit_id="training-code", block=6, excerpt="Policy code | TRAINING-01")])
    if index == 13:
        facts = [(TOPICS[7], 0), (TOPICS[5], 0)]
        question = "Онцгой үед зочин хаана цугларах вэ, ердийн нэвтрэх эрх нь хэдэн цагт дуусах вэ?" if language == "mn" else "Where should a visitor assemble in an emergency, and when does normal visitor access end?"
    elif index == 14:
        facts = [(TOPICS[6], 0), (TOPICS[6], 1)]
        question = "Зардал нэхэмжлэхэд ямар баримт шаардлагатай вэ, нөхөн төлбөрийг хэзээ боловсруулдаг вэ?" if language == "mn" else "Which receipt is required for expenses, and when are reimbursements processed?"
    else:
        return {}
    claims = [topic[3 + fact * 2 + (language == "mn")] for topic, fact in facts]
    return dict(question=question, expected_answer=" ".join(claims), required_claims=claims, evidence=[evidence(topic, fact) for topic, fact in facts])


def generate(directory):
    directory = Path(directory)
    directory.mkdir(parents=True, exist_ok=True)
    sources = {}
    for topic in TOPICS:
        key, en_title, mn_title, en1, mn1, en2, mn2, *_ = topic
        for language, title, facts in (("en", en_title, [en1, en2]), ("mn", mn_title, [mn1, mn2])):
            file = directory / f"{key}-{language}.docx"
            doc = Document()
            doc.add_heading(title, 0)
            doc.add_paragraph("SYNTHETIC LEARNING FIXTURE — NOT COMPANY POLICY" if language == "en" else "СУРГАЛТЫН ЗОХИОМОЛ БАРИМТ — БАЙГУУЛЛАГЫН ЖУРАМ БИШ")
            doc.add_heading("Rules" if language == "en" else "Дүрэм", 1)
            for fact in facts:
                doc.add_paragraph(fact)
            table = doc.add_table(rows=1, cols=2)
            table.rows[0].cells[0].text = "Policy code"
            table.rows[0].cells[1].text = key.upper() + "-01"
            doc.save(file)
            sources[(key, language)] = dict(file=file.name, sha256=hashlib.sha256(file.read_bytes()).hexdigest())
    questions = []
    for question_language, evidence_language in (("mn", "mn"), ("en", "en"), ("mn", "en"), ("en", "mn")):
        direction = f"{question_language}->{evidence_language}"
        for index in range(20):
            answerable = index < 15
            q = dict(id=f"{question_language}-{evidence_language}-{index+1:02}", revision="synthetic-v1", language=question_language,
                evidence_language=evidence_language, direction=direction, split="dev" if index % 2 == 0 else "test", answerable=answerable,
                family=f"fact-{index}", review_status="draft", evidence=[])
            if answerable:
                topic = TOPICS[index % 8]
                fact_index = 0 if index < 8 else 1
                q["question"] = topic[7 + fact_index * 2 + (question_language == "mn")]
                expected = topic[3 + fact_index * 2 + (question_language == "mn")]
                excerpt = topic[3 + fact_index * 2 + (evidence_language == "mn")]
                q.update(expected_answer=expected, required_claims=[expected], evidence=[dict(sources[(topic[0], evidence_language)],
                    unit_id=f"{topic[0]}-{fact_index}", block=4 + fact_index, excerpt=excerpt)])
            else:
                english = ["What is the company pension contribution rate?", "How many paid parental leave weeks are provided?", "What is the CEO's private telephone number?", "Does training cover overseas flights?", "What is the employee stock purchase discount?"]
                mongolian = ["Тэтгэврийн шимтгэлийн хувь хэд вэ?", "Эцэг эхийн цалинтай чөлөө хэдэн долоо хоног вэ?", "Гүйцэтгэх захирлын хувийн утасны дугаар хэд вэ?", "Сургалтын тэтгэмж гадаад нислэгийг хамрах уу?", "Ажилтны хувьцааны хөнгөлөлт хэд вэ?"]
                q.update(question=(mongolian if question_language == "mn" else english)[index - 15], expected_answer="Insufficient evidence", required_claims=[])
            q.update(richer_case(index, question_language, evidence_language, sources))
            questions.append(q)
    dataset = directory / "questions.jsonl"
    dataset.write_text("\n".join(json.dumps(q, ensure_ascii=False) for q in questions) + "\n", encoding="utf8")
    return dataset
