"""Ordered native extraction, optional per-page OCR, and bounded DOC conversion."""
import json
import os
import hashlib
from importlib import metadata as package_metadata
import shutil
import re
import subprocess
import sys
import tempfile
import time
import unicodedata
import zipfile
from pathlib import Path
from pypdf import PdfReader
from docx import Document as WordDocument
from docx.table import Table
from docx.text.paragraph import Paragraph

PARSER_REVISION = "ordered-blocks-v1"


def parser_info(extension):
    """Report parser stack and installed Python package versions, when available."""
    stacks = {
        "pdf": ("pypdf + pypdfium2 + pytesseract", ("pypdf", "pypdfium2", "pytesseract")),
        "docx": ("python-docx", ("python-docx",)),
        "doc": ("LibreOffice conversion + python-docx", ("python-docx",)),
    }
    name, packages = stacks.get(extension, ("unsupported", ()))
    versions = {}
    for package in packages:
        try:
            versions[package] = package_metadata.version(package)
        except package_metadata.PackageNotFoundError:
            versions[package] = None
    return {"name": name, "versions": versions, "revision": PARSER_REVISION}


def classify_error(error, extension=None):
    text = str(error or "").casefold()
    if "encrypted" in text or "password" in text:
        return "encrypted"
    if "no usable text" in text or "empty" in text:
        return "empty_document"
    if "timed out" in text or "deadline" in text or "timeout" in text:
        return "timeout"
    if "upload a pdf" in text or "unsupported" in text or "not supported" in text:
        return "unsupported_format"
    if "ocr" in text or "tesseract" in text:
        return "ocr_failure"
    if extension == "docx" or ("docx" in text and any(word in text for word in ("archive", "zip", "word/document.xml", "invalid"))):
        return "malformed_docx"
    if extension == "pdf" or ("pdf" in text and any(word in text for word in ("invalid", "malformed", "eof", "stream"))):
        return "malformed_pdf"
    if extension == "doc":
        return "malformed_doc_or_conversion_failure"
    return "parse_failure"


def native_fingerprint(config):
    """Invalidate extraction caches when converter/OCR binaries or language data change."""
    result = {}
    for key in ("TESSERACT_CMD", "LIBREOFFICE_CMD"):
        executable = Path(shutil.which(config[key]) or config[key])
        result[key] = hashlib.sha256(executable.read_bytes()).hexdigest() if executable.is_file() else "unavailable"
    tess = Path(config.get("TESSDATA_PREFIX") or (Path(shutil.which(config["TESSERACT_CMD"]) or config["TESSERACT_CMD"]).parent / "tessdata"))
    for language in config["OCR_LANGUAGES"].split("+"):
        file = tess / f"{language}.traineddata"
        result[language] = hashlib.sha256(file.read_bytes()).hexdigest() if file.is_file() else "unavailable"
    return result


def normalize(text):
    text = unicodedata.normalize("NFC", text.replace("\x00", ""))
    return "\n".join(re.sub(r"[^\S\n]+", " ", line).strip() for line in text.splitlines()).strip()


def validate_file(file, extension, config):
    signature = file.read_bytes()[:8]
    if extension == "pdf":
        if not signature.startswith(b"%PDF-"):
            raise ValueError("The file is not a PDF.")
        reader = PdfReader(file)
        if reader.is_encrypted:
            raise ValueError("Encrypted PDFs are not supported. Upload an approved decrypted copy.")
        if len(reader.pages) > config["PDF_MAX_PAGES"]:
            raise ValueError("PDF exceeds the page limit.")
    elif extension == "docx":
        try:
            with zipfile.ZipFile(file) as archive:
                if sum(i.file_size for i in archive.infolist()) > config["DOCX_MAX_EXPANDED"]:
                    raise ValueError("DOCX expanded size exceeds 100 MiB.")
                if "word/document.xml" not in archive.namelist():
                    raise ValueError("DOCX does not contain a Word document.")
                if any("vbaproject" in i.filename.lower() for i in archive.infolist()):
                    raise ValueError("Macro-bearing Word documents are not supported.")
            WordDocument(file)
        except zipfile.BadZipFile as exc:
            raise ValueError("Invalid DOCX archive.") from exc
    elif extension == "doc":
        if signature != bytes.fromhex("d0cf11e0a1b11ae1"):
            raise ValueError("Expected a legacy binary DOC file. Renaming an extension is not conversion.")
    else:
        raise ValueError("Upload a PDF, DOCX, or DOC file.")


def convert_doc(source, destination, config):
    destination.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="rag-lo-") as profile:
        user = Path(profile) / "user"
        user.mkdir()
        # High macro security in an isolated profile; no shared user profile is touched.
        (user / "registrymodifications.xcu").write_text(
            '<?xml version="1.0"?><oor:items xmlns:oor="http://openoffice.org/2001/registry">'
            '<item oor:path="/org.openoffice.Office.Common/Security/Scripting">'
            '<prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop>'
            '</item></oor:items>', encoding="utf8")
        try:
            result = subprocess.run([
                config["LIBREOFFICE_CMD"], f"-env:UserInstallation={Path(profile).as_uri()}",
                "--headless", "--norestore", "--nodefault", "--convert-to", "docx", "--outdir", str(destination), str(source),
            ], capture_output=True, timeout=config["CONVERT_TIMEOUT"], check=False,
                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        except FileNotFoundError as exc:
            raise ValueError("LibreOffice is unavailable. Configure LIBREOFFICE_CMD.") from exc
        except subprocess.TimeoutExpired as exc:
            raise ValueError("DOC conversion timed out.") from exc
        converted = destination / (source.stem + ".docx")
        if result.returncode or not converted.exists():
            raise ValueError("LibreOffice could not convert this DOC file.")
        validate_file(converted, "docx", config)
        return converted


def parse_direct(source, extension, config, forced_pages=()):
    validate_file(source, extension, config)
    warnings, blocks, candidates = [], [], []
    if extension == "doc":
        source = convert_doc(source, source.parent.parent / "converted" / source.stem, config)
        extension = "docx"
        warnings.append("Хуучин legacy DOC форматыг хөрвүүллээ. Хүснэгт болон текстийн дарааллыг шалгана уу.")
    if extension == "docx":
        document = WordDocument(source)
        heading = ""
        for item in document.iter_inner_content():
            kind = "paragraph"
            if isinstance(item, Paragraph):
                text = normalize(item.text)
                if item.style and item.style.name.startswith("Heading"):
                    heading = text
            elif isinstance(item, Table):
                rows = [[normalize(cell.text) for cell in row.cells] for row in item.rows]
                text = "\n".join(" | ".join(row) for row in rows)
                kind = "table"
                warnings.append("Хүснэгтийг мөрөөр танилаа. Нэгтгэсэн нүд болон толгой хэсгийг шалгана уу.")
            else:
                continue
            if text:
                blocks.append(dict(text=text, heading=heading, kind=kind, method="native", block=len(blocks) + 1))
        warnings.append("DOCX байрлалыг хэсэг, гарчгаар тэмдэглэсэн. Текст хайрцаг болон зураг дахь бичвэрийг шалгана уу.")
    else:
        import pypdfium2 as pdfium
        import pytesseract
        pytesseract.pytesseract.tesseract_cmd = config["TESSERACT_CMD"]
        reader = PdfReader(source)
        rendered = None
        try:
            for index, page in enumerate(reader.pages, 1):
                native = normalize(page.extract_text() or "")
                sparse = sum(c.isalnum() for c in native) < config["OCR_MIN_CHARACTERS"]
                corrupt = native.count("\ufffd") / max(1, len(native)) > 0.01
                method, selected, ocr = "native", native, None
                if sparse or corrupt or index in forced_pages:
                    try:
                        if rendered is None:
                            rendered = pdfium.PdfDocument(source)
                        rendered_page = rendered[index - 1]
                        try:
                            width, height = rendered_page.get_size()
                            if width * height * (300 / 72) ** 2 > 25_000_000:
                                raise ValueError("OCR page exceeds the 25 megapixel rendering limit.")
                            bitmap = rendered_page.render(scale=300 / 72)
                            image = bitmap.to_pil()
                            previous_data = os.environ.get("TESSDATA_PREFIX")
                            try:
                                if config.get("TESSDATA_PREFIX"):
                                    os.environ["TESSDATA_PREFIX"] = str(Path(config["TESSDATA_PREFIX"]).resolve())
                                ocr = normalize(pytesseract.image_to_string(image, lang=config["OCR_LANGUAGES"], timeout=config["OCR_TIMEOUT"]))
                            finally:
                                if previous_data is None:
                                    os.environ.pop("TESSDATA_PREFIX", None)
                                else:
                                    os.environ["TESSDATA_PREFIX"] = previous_data
                                image.close()
                                bitmap.close()
                        finally:
                            rendered_page.close()
                    except Exception as exc:
                        raise ValueError(f"OCR failed on page {index}. Check Tesseract, mon/eng data, and page timeout.") from exc
                    method, selected = "ocr", ocr
                    warnings.append(f"{index}-р хуудас: OCR танилтыг ашиглалаа. Тоо, үгүйсгэл, Ө болон Ү үсгийг эхтэй нь тулгана уу.")
                candidates.append(dict(page=index, native=native, ocr=ocr, selected=method))
                blocks.append(dict(text=selected, heading="", kind="page", method=method, page=index, block=index))
        finally:
            if rendered:
                rendered.close()
    if not any(block["text"].strip() for block in blocks):
        raise ValueError("No usable text extracted. Review the source or OCR setup.")
    return dict(blocks=blocks, candidates=candidates, warnings=list(dict.fromkeys(warnings)), parser=PARSER_REVISION)


def parse_bounded(source, extension, config, forced_pages=(), metrics_callback=None):
    """Killable process bounds PDF parsing/rendering too, not only Tesseract."""
    with tempfile.TemporaryDirectory(prefix="rag-parse-") as temporary:
        request = Path(temporary) / "request.json"
        output = Path(temporary) / "output.json"
        metrics = Path(temporary) / "metrics.json"
        keys = ("PDF_MAX_PAGES", "DOCX_MAX_EXPANDED", "OCR_MIN_CHARACTERS", "OCR_TIMEOUT", "CONVERT_TIMEOUT", "TESSERACT_CMD", "TESSDATA_PREFIX", "OCR_LANGUAGES", "LIBREOFFICE_CMD")
        request.write_text(json.dumps(dict(source=str(source), extension=extension, config={key: config[key] for key in keys},
            forced_pages=list(forced_pages), metrics=str(metrics) if metrics_callback else None)), encoding="utf8")
        try:
            result = subprocess.run([sys.executable, "-m", "app.services.parse_worker", str(request), str(output)], capture_output=True,
                timeout=config["PROCESS_TIMEOUT"], creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        except subprocess.TimeoutExpired as exc:
            raise ValueError("Document extraction exceeded the processing deadline.") from exc
        if not output.exists():
            raise ValueError("Document parser process failed. Check the file and native tool installation.")
        data = json.loads(output.read_text(encoding="utf8"))
        if metrics_callback and metrics.exists():
            try:
                metrics_callback(json.loads(metrics.read_text(encoding="utf8")))
            except Exception:
                pass
        if result.returncode or "error" in data:
            raise ValueError(data.get("error", "Document extraction failed."))
        return data
